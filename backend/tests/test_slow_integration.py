# Slow integration (full 3Y pipeline) — pytest -q -m "" backend/tests/test_slow_integration.py
import sys, os, asyncio, pytest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
import main as M
import product_3y as P

PAY = {"product_name":"T","anchor_launch_date":"2029-03-01","horizon_years":2,
  "genre":"Battle Royale","bm_ui":"F2P Cosmetic + Battle Pass","reference_mode":"auto",
  "target_d1":0.50,"organic_share_of_total":0.364,
  "region_mix":{"NA":0.5,"KR":0.5},
  "strategic_hurdle":{"fcy1":100e8,"fcy_start_year_index":1},
  "waves":[{"wave_id":"pc","platform":"PC","offset_months":0,"ua_budget":5e9,"brand_budget":2e9,
            "target_cpa":7500,"d1_nru_normal":0,"prereg_users":1000000,"prereg_activation_rate":0.40}],
  "mode":{"ex_only":{"initial":0.1,"target":0.15,"ramp_days":180},"both":{"initial":0.08,"target":0.25,"ramp_days":180}},
  "synergy":{"retention_lift":1.0,"monetization_lift":1.0}}

def test_official_numbers_frozen_v14_off():
    async def t():
        r0 = await P.run_product_3y(PAY, M.calculate_projection, M.ProjectionInput)
        r3 = await P.run_product_3y({**PAY,"enable_v14_3":True}, M.calculate_projection, M.ProjectionInput)
        assert r0["total"]["gross_krw"] == r3["total"]["gross_krw"]  # A안: preview only
        assert "PREVIEW ONLY" in r3["v14_status"]["v14_3_lifecycle"]["contract"]
        assert r3["v14_status"]["v14_3_lifecycle"]["preview_effect"]["implied_gross_uplift_pct"] >= 0
    asyncio.run(t())

def test_v14_4_owner_gate_blocks():
    async def t():
        with pytest.raises(ValueError):
            await P.run_product_3y({**PAY,"enable_v14_4":True}, M.calculate_projection, M.ProjectionInput)
    asyncio.run(t())

def test_official_scenarios_and_delta_bridge():
    async def t():
        o = await P.run_official_scenarios(PAY, M.calculate_projection, M.ProjectionInput)
        assert o["scenarios"]["worst"]["total"]["gross_krw"] < o["scenarios"]["normal"]["total"]["gross_krw"] < o["scenarios"]["best"]["total"]["gross_krw"]
        assert "not a sales commitment" in o["conditional"]
        d = await P.run_v14_delta_bridge(PAY, M.calculate_projection, M.ProjectionInput)
        assert d["rows"][0]["official"] and any(r["gross_krw"] is None for r in d["rows"])  # prototype Δ 미산출 명시
    asyncio.run(t())

def test_excel_14_sheets():
    async def t():
        r = await P.run_product_3y({**PAY,"enable_bridge":True}, M.calculate_projection, M.ProjectionInput)
        from openpyxl import load_workbook
        from io import BytesIO
        wb = load_workbook(BytesIO(P.build_excel(M.sanitize_for_json(r))))
        assert len(wb.sheetnames) == 14
        assert "12_Projection_Bridge" in wb.sheetnames and "13_Confidence_Badge" in wb.sheetnames
    asyncio.run(t())

def test_envelope_demoted():
    async def t():
        r = await P.run_product_3y(PAY, M.calculate_projection, M.ProjectionInput)
        assert "legacy_lever_envelope" in r and "Do not use for official" in r["legacy_lever_envelope"]["warning"]
        assert "business_scenario_envelope" not in r
    asyncio.run(t())
